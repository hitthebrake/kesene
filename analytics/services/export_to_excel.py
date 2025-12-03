import json
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from django.http import JsonResponse, HttpResponse

from analytics.services.analytics import GetAnalytics


class ExportAnalytics(GetAnalytics):
    def __init__(self, request):
        super().__init__(request)

    @staticmethod
    def write(context):
        wb = Workbook()
        ws = wb.active
        ws.title = context.get('cur_location')
        ws.cell(row=1, column=1).value = "Начальная дата:"
        ws.cell(row=2, column=1).value = "Конечная дата:"
        ws.cell(row=1, column=2).value = context.get('cur_start_date')
        ws.cell(row=2, column=2).value = context.get('cur_end_date')

        cur_row = 4

        for sensor, data in context.get('data', {}).items():
            ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=3)
            ws.cell(row=cur_row, column=1).value = f'{sensor.spot}. {sensor.location}. {sensor.name}. {sensor.sensor_type}'
            cur_column = 1
            max_rows = 0
            for t, values in data.items():
                row = cur_row + 1
                ws.cell(row=row, column=cur_column).value = t[0]
                ws.cell(row=row, column=cur_column + 1).value = "Значение"
                ws.cell(row=row, column=cur_column + 2).value = "Дата"
                row += 1
                max_rows = max(max_rows, len(values.keys()))

                if 'min' in values.keys():
                    ws.cell(row=row, column=cur_column).value = 'Минимальное'
                    try:
                        ws.cell(row=row, column=cur_column + 1).value = f'{float(values['min']['value'])} {t[1]}'
                    except ValueError:
                        ws.cell(row=row, column=cur_column + 1).value = values['min']['value']
                    ws.cell(row=row, column=cur_column + 2).value = values['min']['date']
                    row += 1

                if 'max' in values.keys():
                    ws.cell(row=row, column=cur_column).value = 'Максимальное'
                    try:
                        ws.cell(row=row, column=cur_column + 1).value = f'{float(values['max']['value'])} {t[1]}'
                    except ValueError:
                        ws.cell(row=row, column=cur_column + 1).value = values['max']['value']
                    ws.cell(row=row, column=cur_column + 2).value = values['max']['date']
                    row += 1

                if 'average' in values.keys():
                    ws.cell(row=row, column=cur_column).value = 'Среднее'
                    try:
                        ws.cell(row=row, column=cur_column + 1).value = f'{float(values['average'])} {t[1]}'
                    except ValueError:
                        ws.cell(row=row, column=cur_column + 1).value = values['average']
                cur_column += 4
            cur_row += max_rows + 4

        for col_idx, column in enumerate(ws.columns, start=1):
            col_letter = get_column_letter(col_idx)
            dim = ws.column_dimensions[col_letter]
            current_width = dim.width or 10
            dim.width = current_width * 1.5 if col_idx % 2 == 1 else current_width * 1.5
        return wb

    def export(self):
        context = self.analytics()
        workbook = self.write(context)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", )
        response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(f"Аналитика {context.get('cur_location')}.xlsx")}"
        response["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response["Pragma"] = "no-cache"
        response["Expires"] = "0"

        workbook.save(response)
        return response
