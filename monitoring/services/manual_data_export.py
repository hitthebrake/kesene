import json

from datetime import datetime
from openpyxl import Workbook
from urllib.parse import quote

from django.http import JsonResponse, HttpResponse

from sensors.models import InstanceManualData


class ManualDataExport:
    @staticmethod
    def _parse_date_(date_range):
        start_str = date_range.strip()
        end_str = None
        if "—" in date_range:
            start_str, end_str = [d.strip() for d in date_range.split("—")]

        # Parse each into datetime objects
        start_date = datetime.strptime(start_str, "%d/%m/%y")
        end_date = start_date
        if end_str is not None:
            end_date = datetime.strptime(end_str, "%d/%m/%y")
        return start_date, end_date

    @staticmethod
    def write(data_to_write):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Данные"

        # writing types/date header
        sheet.cell(row=1, column=1).value = "Дата"
        sheet.cell(row=1, column=2).value = "Время"
        sheet.cell(row=1, column=3).value = "Локация"
        sheet.cell(row=1, column=4).value = "Место"
        sheet.cell(row=1, column=5).value = "Специалист"
        sheet.cell(row=1, column=6).value = "Улица.°С"
        sheet.cell(row=1, column=7).value = "Пов.°С"
        sheet.cell(row=1, column=8).value = "Пов.%"
        sheet.cell(row=1, column=9).value = "Vв, м/с"
        for letter in "ABFGHI":
            sheet.column_dimensions[letter].width = 15
        for letter in "CDE":
            sheet.column_dimensions[letter].width = 25

        row = 2
        for data in data_to_write:
            # writing data
            sheet.cell(row=row, column=1).value = data.manual_data.created_date.strftime("%Y-%m-%d")
            sheet.cell(row=row, column=2).value = data.created_time.strftime("%H:%M:%S")
            sheet.cell(row=row, column=3).value = data.manual_data.location
            sheet.cell(row=row, column=4).value = f"№{data.number}. {data.manual_data.spot}"
            sheet.cell(row=row, column=5).value = f"{data.manual_data.user.last_name} {data.manual_data.user.first_name}"
            sheet.cell(row=row, column=6).value = data.manual_data.local_temperature or "#NV"
            sheet.cell(row=row, column=7).value = data.temperature or "#NV"
            sheet.cell(row=row, column=8).value = data.humidity or "#NV"
            sheet.cell(row=row, column=9).value = data.manual_data.wind_speed_in_room or "#NV"
            row += 1
        return workbook

    def export(self, request):
        if request.method == 'POST':
            data = json.loads(request.body)
            date_range = data.get('date-range')
            location = data.get('location')
            spot = data.get('spot')
            start_date, end_date = self._parse_date_(date_range)

            instances = InstanceManualData.objects.select_related('manual_data', 'manual_data__user') \
                .filter(manual_data__is_draft=False, manual_data__created_date__range=(start_date, end_date)) \
                .order_by('manual_data__created_date', 'created_time')

            if location and spot:
                if location == 'Помещения':
                    instances = instances.filter(manual_data__location=spot)
                elif location == 'Периметр' or location == 'Кровля':
                    instances = instances.filter(manual_data__location=location, manual_data__spot=spot)

            workbook = self.write(instances)

            filename = (
                f"Журнал измерении "
                f"{start_date.strftime('%d.%m.%y')}-"
                f"{end_date.strftime('%d.%m.%y')}.xlsx"
            )
            response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", )
            response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
            response["Cache-Control"] = "no-cache, no-store, must-revalidate"
            response["Pragma"] = "no-cache"
            response["Expires"] = "0"

            workbook.save(response)
            return response
        return JsonResponse({"status": "error", 'message':'Произошла ошибка'})