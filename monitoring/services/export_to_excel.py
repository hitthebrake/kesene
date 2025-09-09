import json

from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from urllib.parse import quote

from django.http import JsonResponse, HttpResponse

from webhook.sensor_parsing import sensor_types
from webhook.models import Measurement, Sensor


class ExportToExcel:
    @staticmethod
    def _parse_date_(date_range):
        start_str = date_range.strip()
        end_str = None
        if "—" in date_range:
            start_str, end_str = [d.strip() for d in date_range.split("—")]

        # Parse each into datetime objects
        start_date = datetime.strptime(start_str, "%d/%m/%y")
        end_date = start_date + timedelta(days=1)
        if end_str is not None:
            end_date = datetime.strptime(end_str, "%d/%m/%y")
        return start_date, end_date

    def get_data(self):
        pass

    @staticmethod
    def write(data_to_write):
        workbook = Workbook()
        for sensor_name, data in data_to_write.items():
            sheet = workbook.create_sheet()
            sheet.title = str(sensor_name)[:31]

            # writing types/date header
            sheet.cell(row=1, column=1).value = "Дата/Время"
            sheet.column_dimensions['A'].width = 25
            types = data.get("types")
            for i in range(len(types)):
                sheet.column_dimensions[get_column_letter(i+2)].width = 25
                sheet.cell(row=1, column=i+2).value = f'{types[i][0]} {types[i][1]}'

            # writing data
            row = 2
            objects = data.get("values")
            for obj in objects:
                sheet.cell(row=row, column=1).value = obj.date.strftime("%Y-%m-%d %H:%M:%S")
                sheet.cell(row=row, column=2).value = "#NV" if (obj.value1 is None and len(types)>0) else obj.value1
                sheet.cell(row=row, column=3).value = "#NV" if (obj.value2 is None and len(types)>1) else obj.value2
                sheet.cell(row=row, column=4).value = "#NV" if (obj.value3 is None and len(types)>2) else obj.value3
                row += 1
        del workbook['Sheet']
        return workbook

    def export(self, request):
        if request.method == 'POST':
            data = json.loads(request.body)
            date_range = data.get('date-range')
            location = data.get('location')
            start_date, end_date = self._parse_date_(date_range)

            sensors = Sensor.objects.filter(location=location)
            measurements = Measurement.objects.filter(
                date__range=(start_date, end_date), sensor__in=sensors
            ).order_by('date')

            data_to_write = {}
            for sensor in sensors:
                values = measurements.filter(sensor=sensor)
                sensor_full_name = f'{sensor.name}-{sensor.spot}-{sensor.sensor_type}'
                types = sensor_types[sensor.sensor_type]
                data_to_write[sensor_full_name] = {'types': types, "values": values}

            workbook = self.write(data_to_write)

            response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", )
            response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(f"Отчет {location}.xlsx")}"
            response["Cache-Control"] = "no-cache, no-store, must-revalidate"
            response["Pragma"] = "no-cache"
            response["Expires"] = "0"

            workbook.save(response)
            return response
        return JsonResponse({"status": "error", 'message':'Произошла ошибка'})