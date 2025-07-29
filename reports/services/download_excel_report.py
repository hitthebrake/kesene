from django.http import JsonResponse, HttpResponse
from django.utils.dateparse import parse_date, parse_time

from datetime import datetime
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from urllib.parse import quote

from reports.models import Measurement, Sensor
from reports.extra_functions import *


class DownloadExcelReport:
    stat_labeling = {
        'average': 'Среднее значение',
        'max': 'Максимальное значение',
        'min': 'Минимальное значение',
        'active_duration': 'Активное время',
        'passive_duration': 'Неактивное время',
        'total_count': 'Количество измерении',
        'interval': 'Интервал',
    }
    statistics = {}
    types = {}
    measurements = {}
    location = ""

    def __init__(self, data):
        self.data = data
        self.date_start = parse_date(data.get('date_s'))
        self.time_start = parse_time(data.get('time_s'))
        self.date_end = parse_date(data.get('date_e'))
        self.time_end = parse_time(data.get('time_e'))
        self.sensor_ids = data.get('sensor_ids')



    def download_excel_report(self):
        if not all([self.date_start, self.time_start, self.date_end, self.time_end, self.sensor_ids]):
            return JsonResponse({'status': 'error', 'message': 'Данные некорректно заполнены'})

        start = format_date(datetime.combine(self.date_start, self.time_start))
        end = format_date(datetime.combine(self.date_end, self.time_end))
        sensors = Sensor.objects.filter(id__in=self.sensor_ids)
        data = Measurement.objects.filter(date__range=(start, end), sensor__in=sensors).order_by('date') \
            .values('date', 'value', 'type', 'sensor', 'id')

        for sensor in sensors:
            self.types[sensor.id] = []
            self.statistics[sensor.id] = {}
            for t in sensor.supported_types.all():
                self.statistics[sensor.id][t.id] = calculate_stat(sensor, t, data, end - start)
                self.types[sensor.id].append({
                    'id': t.id,
                    'label': t.label,
                    'code': t.code,
                })

        for d in data:
            if d['sensor'] not in self.measurements.keys():
                self.measurements[d['sensor']] = {}
            if d['date'] not in self.measurements[d['sensor']].keys():
                self.measurements[d['sensor']][d['date']] = {}
            self.measurements[d['sensor']][d['date']][d['type']] = d['value']

        self.location = sensors[0].location if len(sensors) > 0 else 'Место'
        filename = f"Отчет {self.location}.xlsx"
        workbook = self.write(sensors)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", )
        response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
        response["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response["Pragma"] = "no-cache"
        response["Expires"] = "0"

        workbook.save(response)
        return response


    def write(self, sensors):
        workbook = Workbook()
        for sensor in sensors:
            sheet = workbook.create_sheet()
            sheet.title = str(sensor.name)[:31]

            # write types
            col = 2
            sheet.cell(column=1, row=1, value='Дата/Время')
            sheet.column_dimensions['A'].width = 25
            for t in self.types[sensor.id]:
                sheet.cell(column=col, row=1, value=t['label'])
                sheet.column_dimensions[get_column_letter(col)].width = 15
                col += 1

            # writing data
            row = 2
            for date, data in self.measurements[sensor.id].items():
                sheet.cell(column=1, row=row, value=date)
                col = 2
                for t in self.types[sensor.id]:
                    for key, value in data.items():
                        if key == t['id']:
                            sheet.cell(column=col, row=row, value=value)
                    col += 1
                row += 1

            # writing statistics
            col = len(self.statistics[sensor.id]) + 5
            for t in self.types[sensor.id]:
                sheet.column_dimensions[get_column_letter(col)].width = 25
                sheet.column_dimensions[get_column_letter(col + 1)].width = 15
                sheet.cell(column=col, row=3, value=t['label'])
                row = 4
                for key, stat in self.statistics[sensor.id][t['id']].items():
                    sheet.cell(column=col, row=row, value=self.stat_labeling[key])
                    sheet.cell(column=col + 1, row=row, value=stat)
                    row += 1
                col += 3
        del workbook['Sheet']
        return workbook
